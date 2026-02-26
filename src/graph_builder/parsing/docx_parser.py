"""Parser for Microsoft Word .docx files."""

from pathlib import Path

from graph_builder.models import Document
from graph_builder.parsing.utils import SECTION_NUMBER_PATTERN as _SECTION_NUMBER_PATTERN
from graph_builder.parsing.utils import validate_file


class DocxParser:
    """Parser for .docx files using python-docx."""

    SUPPORTED_EXTENSIONS = {".docx"}
    SECTION_NUMBER_PATTERN = _SECTION_NUMBER_PATTERN
    HEADING_STYLES = {"Heading 1", "Heading 2", "Heading 3", "Title"}

    def __init__(self) -> None:
        """Initialize the parser, lazily loading python-docx."""
        self._docx_module = None

    @property
    def docx(self):
        """Lazy-load python-docx module."""
        if self._docx_module is None:
            try:
                import docx

                self._docx_module = docx
            except ImportError as e:
                raise ImportError(
                    "python-docx is required for .docx parsing. "
                    "Install it with: uv sync --extra contracts"
                ) from e
        return self._docx_module

    def supports(self, file_path: Path) -> bool:
        """Check if this parser supports the given file type."""
        return file_path.suffix.lower() in self.SUPPORTED_EXTENSIONS

    def parse(self, file_path: Path) -> Document:
        """Parse a .docx file and return a Document model.

        Args:
            file_path: Path to the .docx file.

        Returns:
            Document with content and paragraph metadata.

        Raises:
            ValueError: If file type is not supported.
            FileNotFoundError: If file does not exist.
        """
        file_path = validate_file(Path(file_path), self.SUPPORTED_EXTENSIONS)

        doc = self.docx.Document(file_path)

        paragraphs_data = []
        content_parts = []
        current_char_offset = 0

        for idx, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if not text:
                continue

            style_name = para.style.name if para.style else "Normal"
            is_heading = style_name in self.HEADING_STYLES

            section_match = _SECTION_NUMBER_PATTERN.match(text)
            section_number = section_match.group(1).rstrip(".") if section_match else None

            para_data = {
                "index": idx,
                "text": text,
                "style": style_name,
                "is_heading": is_heading,
                "section_number": section_number,
                "start_char": current_char_offset,
                "end_char": current_char_offset + len(text),
            }
            paragraphs_data.append(para_data)
            content_parts.append(text)
            current_char_offset += len(text) + 1  # +1 for newline separator

        full_content = "\n".join(content_parts)

        return Document(
            content=full_content,
            source=str(file_path),
            metadata={
                "file_type": "docx",
                "file_name": file_path.name,
                "paragraphs": paragraphs_data,
                "paragraph_count": len(paragraphs_data),
            },
        )
